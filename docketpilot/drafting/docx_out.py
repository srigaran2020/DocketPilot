"""Word output with a removable verification appendix (spec 13)."""
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.shared import Pt

COLOUR_WORD = {"green": "OK", "amber": "CHECK", "red": "PROBLEM"}


def write_docx(path: Path, body: str, report: dict[str, Any], acknowledged: list[str]) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for line in body.splitlines():
        p = doc.add_paragraph()
        # Gaps the papers did not fill stay visible: highlighted so they cannot be filed by mistake.
        for part in re.split(r"(\[(?:missing|NOT IN DOCUMENTS):[^\]]*\])", line, flags=re.I):
            if not part:
                continue
            run = p.add_run(part)
            if re.match(r"\[(?:missing|NOT IN DOCUMENTS):", part, re.I):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        letters = re.sub(r"[^A-Za-z]", "", line)
        if letters and line.strip() == line.strip().upper() and len(line) < 120:  # headings in capitals
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Appendix: DocketPilot verification summary (remove before filing)", level=2)
    doc.add_paragraph(f"Release status at approval: {report.get('release')}. Relevant date: {report.get('relevant_date')}.")
    for c in report.get("citations", []):
        ack = " (acknowledged by advocate)" if c["id"] in acknowledged else ""
        doc.add_paragraph(f"[{COLOUR_WORD.get(c['colour'], c['colour'])}] {c['raw']} - {c['verdict']}: {c['reason']}{ack}",
                          style="List Bullet")
    for s in report.get("statutes", []):
        ack = " (acknowledged by advocate)" if s["id"] in acknowledged else ""
        doc.add_paragraph(f"[{COLOUR_WORD.get(s['colour'], s['colour'])}] {s['raw']} - {s['status']}: {s['message']}{ack}",
                          style="List Bullet")
    for c in report.get("corrections", []):
        doc.add_paragraph(f"[FIXED] {c['original']} -> {c.get('replacement') or 'removed'}: {c['reason']}", style="List Bullet")
    doc.add_paragraph(report.get("notice", ""))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
