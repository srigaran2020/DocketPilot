"""Build the DocketPilot project report from the friend's Word file.

Everything up to and including the Table of Contents heading is kept as it is. The TOC table, abstract, lists of tables
and figures and all chapters are written here, in the paragraph formats of the original (Times New Roman, 14 pt body,
1.15 line spacing, 6 pt after). Word is used to measure where each heading lands, then the page numbers are filled in.

    .venv\\Scripts\\python docs\\report\\build_report.py
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
import content as C  # noqa: E402

SRC = ROOT / "docs" / "DocketPilot_sample.docx"
OUT = HERE / "DocketPilot_Project_Report.docx"
PDF = HERE / "DocketPilot_Project_Report.pdf"
TEXT_W = 9026  # twips between the 1-inch margins of an A4 page
LAST_PAGE = 40  # appendix ends here
LINES_PER_PAGE = 44  # appendix code, Times New Roman 12 pt single spaced (first guess, refined by measuring)


# ---------------------------------------------------------------- low-level helpers
def _set_run(r, size=14, bold=False, italic=False):
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    return r


def _fmt(p, line=1.15, after=6, before=0, align=None, keep=False, page_break=False):
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    pf.keep_with_next = keep
    pf.page_break_before = page_break
    return p


def _cell_margins(tc, m=90):
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(m))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def _table_borders(tbl, color="000000", size=4, none=False):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "none" if none else "single")
        e.set(qn("w:sz"), "0" if none else str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto" if none else color)
        b.append(e)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    look = tblPr.find(qn("w:tblLook"))
    for el in (b, lay):  # schema order: ... tblBorders, shd, tblLayout, tblCellMar, tblLook
        if look is not None:
            look.addprevious(el)
        else:
            tblPr.append(el)


def _widths(tbl, widths):
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths, strict=True):
        gc.set(qn("w:w"), str(w))
    for row in tbl.rows:
        for c, w in zip(row.cells, widths, strict=True):
            tcPr = c._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.insert(0, tcW)
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    trPr.append(e)


def _cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _replace_cell_text(tc_el, text, align=None, bold=None):
    """Put text in a cloned table cell, keeping the run formatting of its first run (or paragraph mark)."""
    ps = tc_el.findall(qn("w:p"))
    for extra in ps[1:]:
        tc_el.remove(extra)
    p = ps[0]
    runs = p.findall(qn("w:r"))
    rpr = None
    if runs and runs[0].find(qn("w:rPr")) is not None:
        rpr = copy.deepcopy(runs[0].find(qn("w:rPr")))
    elif p.find(qn("w:pPr")) is not None and p.find(qn("w:pPr")).find(qn("w:rPr")) is not None:
        rpr = copy.deepcopy(p.find(qn("w:pPr")).find(qn("w:rPr")))
    for r in runs:
        p.remove(r)
    for el in list(p):
        if el.tag != qn("w:pPr"):
            p.remove(el)
    if align is not None:
        ppr = p.find(qn("w:pPr"))
        jc = ppr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            ppr.append(jc)
        jc.set(qn("w:val"), align)
    if not text:
        return
    r = OxmlElement("w:r")
    if rpr is not None:
        if bold is not None:
            for b in rpr.findall(qn("w:b")) + rpr.findall(qn("w:bCs")):
                rpr.remove(b)
            if bold:
                rpr.insert(0, OxmlElement("w:b"))
        r.append(rpr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    p.append(r)


# ---------------------------------------------------------------- document writer
class Report:
    def __init__(self, pages: dict | None, code_lines: int):
        self.d = Document(SRC)
        self.pages = pages or {}
        self.code_lines = code_lines
        self.toc: list[tuple[str, str, str, bool]] = []  # (number, title, key, is_chapter)
        self.figs: list[tuple[str, str]] = []
        self.tabs: list[tuple[str, str]] = []
        self.marks: list[str] = []  # paragraph texts to locate in Word
        self.first_chapter = True
        self.old_lit = [[" ".join(q.text.strip() for q in c.paragraphs if q.text.strip()) for c in r.cells]
                        for r in list(self.d.tables[5].rows)[1:]]

    # ---- blocks
    def chapter(self, num: str | None, title: str):
        if num:
            p = self.d.add_paragraph()
            _fmt(p, line=1.0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, page_break=not self.first_chapter)
            _set_run(p.add_run(f"CHAPTER {num}"), 14)
            self.marks.append(f"CHAPTER {num}")
        p = self.d.add_paragraph()
        _fmt(p, line=1.15, after=6, align=WD_ALIGN_PARAGRAPH.CENTER, page_break=(num is None))
        _set_run(p.add_run(title), 14, bold=True)
        self.first_chapter = False
        blank = self.d.add_paragraph()
        _fmt(blank, line=1.0, after=0)
        key = f"CHAPTER {num}" if num else title
        if not num:
            self.marks.append(title)
        self.toc.append((num or "", title, key, True))

    def heading(self, num: str, title: str):
        p = self.d.add_paragraph()
        _fmt(p, line=1.15, after=6, before=10, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)
        text = f"{num}  {title}" if num else title
        _set_run(p.add_run(text), 14, bold=True)
        self.marks.append(text)
        if num:
            self.toc.append(("", f"{num}  {title}", text, False))

    def para(self, text: str):
        p = self.d.add_paragraph()
        _fmt(p, line=1.15, after=6)
        _set_run(p.add_run(text), 14)

    def bullets(self, items):
        for it in items:
            p = self.d.add_paragraph()
            _fmt(p, line=1.15, after=4)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
            _set_run(p.add_run("•\t" + it), 14)

    def _caption(self, cap: str, container=None, size=12):
        label, rest = re.match(r"((?:Fig\.|Table) \d+\.\d+)\s+(.*)", cap).groups()
        p = container.add_paragraph() if container is not None else self.d.add_paragraph()
        _fmt(p, line=1.15, after=6, before=3, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_run(p.add_run(label + " "), size, bold=True)
        _set_run(p.add_run(rest), size)
        self.marks.append(cap)
        (self.figs if label.startswith("Fig") else self.tabs).append((label.split(" ")[1], rest))
        return p

    def figure(self, path: Path, width: float, cap: str):
        p = self.d.add_paragraph()
        _fmt(p, line=1.0, after=0, before=6, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
        p.add_run().add_picture(str(path), width=Inches(width))
        self._caption(cap)

    def figure_pair(self, items):
        t = self.d.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t, none=True)
        _widths(t, [TEXT_W // 2, TEXT_W - TEXT_W // 2])
        _cant_split(t.rows[0])
        for cell, (path, cap) in zip(t.rows[0].cells, items, strict=True):
            p = cell.paragraphs[0]
            _fmt(p, line=1.0, after=0, before=6, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
            p.add_run().add_picture(str(path), width=Inches(2.85))
            self._caption(cap, container=cell)
        gap = self.d.add_paragraph()
        _fmt(gap, line=1.0, after=0)

    def table(self, cap: str, header, rows, widths, size=12):
        cp = self.d.add_paragraph()
        _fmt(cp, line=1.15, after=4, before=8, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
        label, rest = re.match(r"(Table \d+\.\d+)\s+(.*)", cap).groups()
        _set_run(cp.add_run(label + " "), size, bold=True)
        _set_run(cp.add_run(rest), size, bold=True)
        self.marks.append(cap)
        self.tabs.append((label.split(" ")[1], rest))
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t)
        _widths(t, widths)
        for i, row in enumerate([header] + rows):
            _cant_split(t.rows[i])
            for j, (cell, txt) in enumerate(zip(t.rows[i].cells, row, strict=True)):
                p = cell.paragraphs[0]
                _fmt(p, line=1.0, after=2, before=2,
                     align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 or j > 0 and len(str(txt)) < 14 else WD_ALIGN_PARAGRAPH.LEFT)
                _set_run(p.add_run(str(txt)), size, bold=(i == 0))
        _repeat_header(t.rows[0])
        gap = self.d.add_paragraph()
        _fmt(gap, line=1.0, after=0)

    def lit_table(self, rows):
        """Literature survey in the original's cell format (14 pt, 90-twip margins), resized to the page width."""
        header = ["S. No.", "Title of the Paper", "Name of the Journal / Source", "Year",
                  "Algorithms / Techniques", "Findings"]
        widths = [620, 1800, 1560, 800, 2040, 2206]  # sums to the text width; no word breaks at 14 pt
        t = self.d.add_table(rows=1 + len(rows), cols=6)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t)
        _widths(t, widths)
        for i, row in enumerate([header] + rows):
            for cell, txt in zip(t.rows[i].cells, row, strict=True):
                _cell_margins(cell._tc)
                p = cell.paragraphs[0]
                _fmt(p, line=1.0, after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
                _set_run(p.add_run(str(txt)), 14, bold=(i == 0))
        _cant_split(t.rows[0])
        _repeat_header(t.rows[0])

    def code(self, lines):
        for ln in lines:
            p = self.d.add_paragraph()
            _fmt(p, line=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.left_indent = Inches(0)
            _set_run(p.add_run(ln.rstrip().replace("\t", "    ") or " "), 12)

    def blocks(self, items):
        for b in items:
            k = b[0]
            if k == "p":
                self.para(b[1])
            elif k == "h":
                self.heading("", b[1])
            elif k == "bul":
                self.bullets(b[1])
            elif k == "fig":
                self.figure(b[1], b[2], b[3])
            elif k == "fig2":
                self.figure_pair(b[1])
            elif k == "table":
                self.table(b[1], b[2], b[3], b[4])

    # ---- front matter (after the TOC heading)
    def front(self):
        els = list(self.d.element.body.iterchildren())
        # TOC table rows
        toc = self.d.tables[2]
        chap_tpl, sub_tpl = copy.deepcopy(toc.rows[4]._tr), copy.deepcopy(toc.rows[8]._tr)
        for r in list(toc.rows)[4:]:
            toc._tbl.remove(r._tr)
        for num, title, key, is_ch in self.toc:
            tr = copy.deepcopy(chap_tpl if is_ch else sub_tpl)
            tcs = tr.findall(qn("w:tc"))
            _replace_cell_text(tcs[0], num)
            _replace_cell_text(tcs[1], title, align="center" if is_ch else None)
            _replace_cell_text(tcs[2], str(self.pages.get(key, "")))
            trh = tr.find(qn("w:trPr")).find(qn("w:trHeight"))
            trh.set(qn("w:val"), "520" if is_ch else "400")
            toc._tbl.append(tr)
        # abstract text and keywords
        abs_p, kw_p = els[81], els[84]
        runs = abs_p.findall(qn("w:r"))
        runs[0].find(qn("w:t")).text = C.ABSTRACT
        for r in runs[1:]:
            abs_p.remove(r)
        kr = kw_p.findall(qn("w:r"))
        kt = kr[2].find(qn("w:t"))
        kt.text = " " + C.KEYWORDS
        kt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for r in kr[3:]:
            kw_p.remove(r)
        # the blank lines that pushed LIST OF TABLES to the next page -> a page break instead
        for el in els[85:112]:
            el.getparent().remove(el)
        self._page_break_before(els[112])
        for el in els[115:119]:  # blank lines between the two lists
            el.getparent().remove(el)
        self._page_break_before(els[119])
        for tbl, items, prefix in ((self.d.tables[3], self.tabs, "Table "), (self.d.tables[4], self.figs, "Fig. ")):
            tpl = copy.deepcopy(tbl.rows[1]._tr)
            for r in list(tbl.rows)[1:]:
                tbl._tbl.remove(r._tr)
            for no, name in items:
                tr = copy.deepcopy(tpl)
                tcs = tr.findall(qn("w:tc"))
                _replace_cell_text(tcs[0], no, bold=False)
                _replace_cell_text(tcs[1], name, align="left", bold=False)
                key = prefix + no + " " + name
                _replace_cell_text(tcs[2], str(self.pages.get(key, "")), bold=False)
                tbl._tbl.append(tr)

    @staticmethod
    def _page_break_before(p_el):
        ppr = p_el.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p_el.insert(0, ppr)
        ppr.insert(0, OxmlElement("w:pageBreakBefore"))

    # ---- body
    def body(self):
        els = list(self.d.element.body.iterchildren())
        for el in els[125:-1]:  # friend's chapter text; keep the bookmark end (124) and the final section
            el.getparent().remove(el)

        self.chapter("1", "INTRODUCTION")
        for t in C.INTRO:
            self.para(t)

        self.chapter("2", "LITERATURE SURVEY")
        self.para(C.LIT_INTRO)
        self.lit_table(self._lit_rows())
        for t in C.LIT_GAP:
            self.para(t)
        self.table(*C.LIT_TABLE)

        self.chapter("3", "METHODOLOGY")
        self.blocks(C.M["3"])
        for key, title in (("3.1", "MODEL ARCHITECTURE"), ("3.2", "DATA COLLECTION AND PREPARATION"),
                           ("3.3", "PRIVACY GATEWAY AND DOCUMENT INTAKE"), ("3.4", "MULTI-AGENT WORKFLOW"),
                           ("3.5", "HYBRID RETRIEVAL AND RE-RANKING"),
                           ("3.6", "CITATION VERIFICATION AND AUTO-CORRECTION"), ("3.7", "EVALUATION METHODOLOGY")):
            self.heading(key, title)
            self.blocks(C.M[key])

        self.chapter("4", "RESULTS AND DISCUSSION")
        self.blocks(C.R["4"])
        for key, title in (("4.1", "APPLICATION INTERFACE"), ("4.2", "EFFECT OF EACH PIPELINE LAYER"),
                           ("4.3", "CITATION VERIFICATION ACCURACY"), ("4.4", "RETRIEVAL AND PRIVACY RESULTS"),
                           ("4.5", "COST AND SPEED"), ("4.6", "COMPARISON WITH EXISTING SYSTEMS"),
                           ("4.7", "EXPLAINABILITY WITH LIME")):
            self.heading(key, title)
            self.blocks(C.R[key])

        self.chapter("5", "CONCLUSION AND FUTURE WORK")
        self.blocks(C.CONCLUSION)
        self.marks.append("END_OF_CONCLUSION")
        end = self.d.add_paragraph()
        _fmt(end, line=1.0, after=0)

        self.chapter(None, "REFERENCES")
        self.toc[-1] = ("6", "REFERENCES", "REFERENCES", True)
        for i, ref in enumerate(C.REFERENCES, 1):
            p = self.d.add_paragraph()
            _fmt(p, line=1.15, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.45)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
            _set_run(p.add_run(f"[{i}]\t{ref}"), 14)

        self.chapter(None, "APPENDIX")
        self.toc[-1] = ("7", "APPENDIX", "APPENDIX", True)
        budget = self.code_lines
        for path, title in C.APPENDIX_FILES:
            if budget <= 0:
                break
            lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
            take = lines[:budget]
            p = self.d.add_paragraph()
            _fmt(p, line=1.15, after=4, before=8, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)
            _set_run(p.add_run(f"{title}  ({path})"), 14, bold=True)
            self.code(take)
            budget -= len(take) + 2
        self.marks.append("END_OF_APPENDIX")

    def _lit_rows(self):
        """Rows 1-5 from the friend's table (wording cleaned), then the new ones."""
        rows = self.old_lit
        fixes = [("The PPT notes that its focus is", "Its stated focus is"),
                 ("The PPT summarizes the review as finding", "The review finds"),
                 (" Full bibliographic details were not confirmed in the provided material.", ""),
                 (" (as listed in the project PPT)", ""), ("important challen.", "important challenges."),
                 ("Retrieval- Augmented", "Retrieval-Augmented"), ("AI- Driven", "AI-Driven"),
                 ("AI- assisted", "AI-assisted"), ("document- processing", "document-processing"),
                 ("personal- information", "personal-information")]
        out = []
        for row in rows:
            row = [re.sub(r"\s+", " ", c) for c in row]
            for a, b in fixes:
                row = [c.replace(a, b) for c in row]
            out.append(row)
        for i, r in enumerate(C.LIT_NEW, len(out) + 1):
            out.append([str(i)] + r)
        return out

    def footer(self):
        """The original page number sat in a one-digit-wide text box; use a centred PAGE field instead."""
        ftr = self.d.sections[-1].footer
        for p in list(ftr._element):
            ftr._element.remove(p)
        p = ftr.add_paragraph()
        _fmt(p, line=1.0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
            r = _set_run(p.add_run(), 12)
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r._r.append(fc)
            elif text == " PAGE ":
                it = OxmlElement("w:instrText")
                it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                it.text = text
                r._r.append(it)
            else:
                r.text = text

    def save(self):
        self.body()
        self.front()
        self.footer()
        self.d.save(OUT)


# ---------------------------------------------------------------- measuring in Word
def measure(path: Path, marks: list[str]) -> tuple[dict, int, int]:
    import win32com.client as win

    word = win.DispatchEx("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(str(path), False, True)
    try:
        doc.Repaginate()
        sec = doc.Sections(doc.Sections.Count)
        pages, conc_end = {}, 0
        want = [m for m in marks if not m.startswith("END_")]
        last_text_page = prev_pg = 0
        for p in sec.Range.Paragraphs:
            t = p.Range.Text.strip().replace("\x07", "")
            if not t:
                continue
            pg = p.Range.Information(1)
            last_text_page = pg
            for m in want:
                if m not in pages and t.startswith(m[:60]):
                    pages[m] = pg
            if t == "REFERENCES" and not conc_end:
                conc_end = prev_pg
            prev_pg = pg
        doc.SaveAs2(str(PDF), 17)
        return pages, conc_end, last_text_page
    finally:
        doc.Close(False)
        word.Quit()


def main():
    code_lines = int(sys.argv[1]) if len(sys.argv) > 1 else 10 * LINES_PER_PAGE
    pages = None
    for attempt in range(6):
        rep = Report(pages, code_lines)
        rep.save()
        found, conc_end, last = measure(OUT, rep.marks)
        print(f"attempt {attempt}: code lines {code_lines}, conclusion ends p.{conc_end}, appendix ends p.{last}")
        if last == LAST_PAGE and pages is not None:
            break
        if last != LAST_PAGE:
            code_lines += (LAST_PAGE - last) * LINES_PER_PAGE - (6 if last > LAST_PAGE else 0)
        pages = {**found}
        pages.update({f"Fig. {no} {name}": found.get(f"Fig. {no} {name}", "") for no, name in rep.figs})
        pages.update({f"Table {no} {name}": found.get(f"Table {no} {name}", "") for no, name in rep.tabs})
    print("wrote", OUT, "and", PDF)
    missing = [m for m in rep.marks if not m.startswith("END_") and m not in found]
    if missing:
        print("not located:", missing[:10])


if __name__ == "__main__":
    main()
